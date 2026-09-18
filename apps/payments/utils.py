import re
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree


def _parse_labeled_text(text: str):
    return parse_single_mtn_text(text)


def _normalise_header(value):
    return re.sub(r'[^a-z0-9]', '', str(value or '').lower())


def _parse_spreadsheet_row(headers, values):
    aliases = {
        'transaction_id': {'txnid', 'transactionid', 'transaction', 'transid', 'id'},
        'payment_date': {'date', 'paymentdate'},
        'payment_time': {'time', 'paymenttime'},
        'sender_name': {'from', 'sender', 'sendername', 'name', 'fromhandlername', 'senderhandlername', 'payername'},
        'sender_phone': {'tel', 'phone', 'senderphone', 'telephone'},
        'amount': {'amount', 'paymentamount', 'amountpaid'},
    }
    data = {key: '' for key in aliases}
    for header, value in zip(headers, values):
        normalised = _normalise_header(header)
        target = next((key for key, names in aliases.items() if normalised in names), None)
        if target:
            data[target] = value

    text = '\n'.join(f'{key}: {value}' for key, value in data.items() if value not in ('', None))
    parsed = parse_single_mtn_text(text)
    parsed['transaction_id'] = str(data['transaction_id'] or '').strip()
    parsed['sender_name'] = str(data['sender_name'] or '').strip()
    parsed['sender_phone'] = str(data['sender_phone'] or '').strip()
    if data['amount'] not in ('', None):
        parsed['amount'] = Decimal(str(data['amount']).replace(',', '').replace('UGX', '').strip())
    if data['payment_date']:
        if hasattr(data['payment_date'], 'year'):
            parsed['payment_date'] = data['payment_date']
        else:
            # Word exports commonly combine date and time in one cell, e.g.
            # "08/03/2026, 10:25:06".
            date_time = _parse_date_time(str(data['payment_date']))
            if date_time:
                parsed['payment_date'] = date_time.date()
                parsed['payment_time'] = date_time.time()
    if data['payment_time']:
        parsed['payment_time'] = data['payment_time'] if hasattr(data['payment_time'], 'hour') else parsed['payment_time']
    return parsed


def _parse_date_time(value):
    value = value.strip().replace('\xa0', ' ')
    for date_format in ('%d/%m/%Y, %H:%M:%S', '%d/%m/%Y %H:%M:%S', '%d/%m/%Y, %H:%M', '%d/%m/%Y %H:%M'):
        try:
            return datetime.strptime(value, date_format)
        except ValueError:
            continue
    return None


def parse_uploaded_payments(uploaded_file):
    filename = uploaded_file.name.lower()
    content = uploaded_file.read()
    if filename.endswith('.xlsx'):
        from openpyxl import load_workbook

        workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = rows[0]
        return [_parse_spreadsheet_row(headers, row) for row in rows[1:] if any(value not in (None, '') for value in row)]

    if filename.endswith('.docx'):
        namespace = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        with ZipFile(BytesIO(content)) as archive:
            document = ElementTree.fromstring(archive.read('word/document.xml'))
        paragraphs = []
        for paragraph in document.findall('.//w:body/w:p', namespace):
            text = ''.join(node.text or '' for node in paragraph.findall('.//w:t', namespace))
            if text.strip():
                paragraphs.append(text)
        paragraph_text = '\n'.join(paragraphs)
        parsed = parse_mtn_texts(paragraph_text)
        for table in document.findall('.//w:tbl', namespace):
            rows = table.findall('./w:tr', namespace)
            if not rows:
                continue
            headers = [
                ''.join(node.text or '' for node in cell.findall('.//w:t', namespace))
                for cell in rows[0].findall('./w:tc', namespace)
            ]
            parsed.extend(
                _parse_spreadsheet_row(
                    headers,
                    [''.join(node.text or '' for node in cell.findall('.//w:t', namespace)) for cell in row.findall('./w:tc', namespace)],
                )
                for row in rows[1:]
            )
        return parsed

    raise ValueError('Unsupported file type. Upload a .docx or .xlsx file.')


def parse_single_mtn_text(text: str):
    data = {
        'transaction_id': '',
        'payment_date': None,
        'payment_time': None,
        'sender_name': '',
        'sender_phone': '',
        'amount': Decimal('0'),
        'raw_text': text,
    }

    txn_match = re.search(r'TxnID:\s*([^\s]+)', text, re.I)
    if txn_match:
        data['transaction_id'] = txn_match.group(1).strip()

    date_match = re.search(r'Date:\s*(\d{2}/\d{2}/\d{4})', text)
    if date_match:
        data['payment_date'] = datetime.strptime(date_match.group(1), '%d/%m/%Y').date()

    time_match = re.search(r'Time:\s*(\d{1,2}:\d{2})', text)
    if time_match:
        data['payment_time'] = datetime.strptime(time_match.group(1), '%H:%M').time()

    sender_match = re.search(r'From:\s*([^\n]+)', text, re.I)
    if sender_match:
        data['sender_name'] = sender_match.group(1).strip()

    phone_match = re.search(r'Tel:\s*([0-9+\s]+)', text, re.I)
    if phone_match:
        data['sender_phone'] = phone_match.group(1).strip()

    amount_match = re.search(r'Amount:\s*UGX\s*([0-9,\.]+)', text, re.I)
    if amount_match:
        data['amount'] = Decimal(amount_match.group(1).replace(',', ''))

    return data


def parse_mtn_text(raw_text: str):
    payments = parse_mtn_texts(raw_text)
    return payments[0] if payments else parse_single_mtn_text(raw_text)


def parse_mtn_texts(raw_text: str):
    text = raw_text or ''
    blocks = re.split(r'(?=TxnID:\s*)', text.strip())
    parsed = []
    for block in blocks:
        block = block.strip()
        if not block:
            continue
        parsed.append(parse_single_mtn_text(block))
    if not parsed and text:
        parsed.append(parse_single_mtn_text(text))
    return parsed
